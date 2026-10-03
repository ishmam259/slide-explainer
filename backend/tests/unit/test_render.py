"""All four renderers produce the same sections; HTML escapes model text; math renders."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from slidex.core.hardware import chromium_installed
from slidex.db.files import FileStore
from slidex.llm.schemas import (
    DocSection,
    EvidenceRef,
    ExplanationBlock,
    ExplanationDocumentModel,
    FormulaBlock,
    ReconstructedBlock,
    ReferBackBlock,
    SlideSaysBlock,
    SourceRef,
)
from slidex.render import RENDERERS, available_formats

needs_browser = pytest.mark.skipif(not chromium_installed(), reason="Chromium not installed")

EV = {
    "S1": EvidenceRef(kind="slide", label="Slide 1", slide=1),
    "W1": EvidenceRef(kind="web", label="Quorum docs › Overlap", url="https://docs.example.org/q"),
}


def _doc() -> ExplanationDocumentModel:
    return ExplanationDocumentModel(
        title="Replication",
        organization="by_slide",
        sections=[
            DocSection(
                slide_numbers=[1],
                title="Slide 1: Quorums",
                evidence=EV,
                blocks=[
                    SlideSaysBlock(text="R + W > N", citations=["S1"]),
                    ExplanationBlock(
                        markdown="Reads overlap writes. <script>alert(1)</script>", citations=["W1"]
                    ),
                    FormulaBlock(
                        latex=r"R + W > N",
                        steps_markdown="Each read meets a write.",
                        citations=["W1"],
                    ),
                ],
            ),
            DocSection(
                slide_numbers=[2],
                title="Slide 2: Vague",
                evidence=EV,
                blocks=[
                    ReconstructedBlock(
                        markdown="Likely about sloppy quorums.", confidence=0.6, citations=["W1"]
                    ),
                    ReferBackBlock(concept="quorum", slide=1),
                ],
            ),
        ],
        sources=[
            SourceRef(
                source_id="s",
                title="Quorum docs",
                kind="documentation",
                url="https://docs.example.org/q",
            )
        ],
    )


def test_all_formats_registered() -> None:
    assert available_formats() >= {"md", "html", "pdf", "docx"}


def test_markdown_zip(tmp_path: Path) -> None:
    fn, ext = RENDERERS["md"]
    data = fn(_doc(), FileStore(tmp_path))
    text = zipfile.ZipFile(io.BytesIO(data)).read("document.md").decode()
    assert ext == "zip" and text.index("Slide 1: Quorums") < text.index("Slide 2: Vague")
    assert "$$\nR + W > N\n$$" in text and "confidence 60%" in text


@needs_browser
def test_html_escapes_and_renders_math(tmp_path: Path) -> None:
    html = RENDERERS["html"][0](_doc(), FileStore(tmp_path)).decode()
    assert "<script>alert(1)</script>" not in html  # model text is escaped
    assert "katex" in html and "Quorum docs › Overlap" in html
    assert html.index("Slide 1: Quorums") < html.index("Slide 2: Vague")


@needs_browser
def test_pdf_has_pages_text_and_links(tmp_path: Path) -> None:
    pdf = RENDERERS["pdf"][0](_doc(), FileStore(tmp_path))
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        text = "".join(str(p.get_text()) for p in doc.pages())
        links = [lk for p in doc.pages() for lk in p.get_links()]
    assert "Slide 1: Quorums" in text and "Slide 2: Vague" in text
    assert any(lk.get("uri") == "https://docs.example.org/q" for lk in links)


@needs_browser
def test_docx_headings_formula_image_and_sources(tmp_path: Path) -> None:
    data = RENDERERS["docx"][0](_doc(), FileStore(tmp_path))
    d = Document(io.BytesIO(data))
    headings = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["Slide 1: Quorums", "Slide 2: Vague", "Sources"]
    assert len(d.inline_shapes) == 1  # the formula image
    assert any("see Slide 1" in p.text for p in d.paragraphs)
