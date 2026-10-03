"""Intake: PPTX/PDF/images → slides (fixtures generated in tmp dirs; LibreOffice mocked)."""

from __future__ import annotations

import io
import subprocess
import zipfile
from pathlib import Path
from typing import Any

import pymupdf
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from slidex.core.config import get_settings
from slidex.core.errors import SlidexError
from slidex.db.files import FileStore
from slidex.intake import pptx as pptx_mod
from slidex.intake.common import validate_upload
from slidex.intake.images import normalize_image, read_images
from slidex.intake.pdf import read_pdf


def make_pptx(path: Path) -> Path:
    pres = Presentation()
    s1 = pres.slides.add_slide(pres.slide_layouts[5])
    s1.shapes.title.text = "Replication"
    box = s1.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
    box.text_frame.text = "Copies of data on several nodes"
    s1.notes_slide.notes_text_frame.text = "Mention leader election."
    s2 = pres.slides.add_slide(pres.slide_layouts[5])
    s2.shapes.title.text = "Quorum table"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "R", "W"
    table.cell(1, 0).text, table.cell(1, 1).text = "2", "2"
    hidden = pres.slides.add_slide(pres.slide_layouts[5])
    hidden.shapes.title.text = "Hidden"
    hidden._element.set("show", "0")
    pres.save(str(path))
    return path


def make_pdf(path: Path, pages: list[str]) -> Path:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text)
    doc.save(str(path))
    return path


def test_pdf_text_render_and_image_only(tmp_path: Path) -> None:
    pdf = make_pdf(tmp_path / "d.pdf", ["Consistent hashing places keys on a ring", ""])
    slides = read_pdf(pdf, FileStore(tmp_path / "files"))
    assert [s.number for s in slides] == [1, 2]
    assert "Consistent hashing" in slides[0].native_text and not slides[0].image_only
    assert slides[1].image_only
    assert len(slides[0].image_hash) == 64


def test_pdf_corrupt(tmp_path: Path) -> None:
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"not a pdf")
    with pytest.raises(SlidexError) as exc:
        read_pdf(bad, FileStore(tmp_path / "f"))
    assert exc.value.code == "corrupt_file"


def test_pptx_text_notes_tables_hidden(tmp_path: Path) -> None:
    content = pptx_mod.extract_pptx(make_pptx(tmp_path / "d.pptx"))
    assert len(content) == 2  # hidden slide skipped
    text, notes, _tables = content[0]
    assert text.splitlines()[0] == "Replication"  # reading order: title first
    assert notes == "Mention leader election."
    assert content[1][2] == [[["R", "W"], ["2", "2"]]]


def test_pptx_render_uses_headless_libreoffice(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    deck = make_pptx(tmp_path / "d.pptx")
    calls: list[list[str]] = []

    def fake_run(cmd: list[str], **_: Any) -> subprocess.CompletedProcess[bytes]:
        calls.append(cmd)
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        make_pdf(outdir / "d.pdf", ["one", "two"])
        return subprocess.CompletedProcess(cmd, 0)

    monkeypatch.setattr(pptx_mod.subprocess, "run", fake_run)
    monkeypatch.setattr(pptx_mod, "libreoffice_path", lambda _: Path("soffice.exe"))
    slides = pptx_mod.read_pptx(deck, FileStore(tmp_path / "files"), get_settings())
    assert calls and "--headless" in calls[0] and "pdf" in calls[0]
    assert [s.notes for s in slides] == ["Mention leader election.", ""]


def test_pptx_without_libreoffice(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(pptx_mod, "libreoffice_path", lambda _: None)
    with pytest.raises(SlidexError) as exc:
        pptx_mod.render_pptx(make_pptx(tmp_path / "d.pptx"), get_settings())
    assert exc.value.code == "libreoffice_missing"


def test_images_exif_rgb_and_resize(tmp_path: Path) -> None:
    img = Image.new("RGBA", (4000, 1000), (255, 0, 0, 128))
    exif = img.getexif()
    exif[0x0112] = 6  # rotate 90°
    buf = io.BytesIO()
    img.convert("RGB").save(buf, "JPEG", exif=exif)
    out = Image.open(io.BytesIO(normalize_image(buf.getvalue())))
    assert out.mode == "RGB" and max(out.size) == 2000
    assert out.size[1] > out.size[0]  # rotated to portrait
    p1 = tmp_path / "b.png"
    p2 = tmp_path / "a.png"
    Image.new("RGB", (10, 10)).save(p1)
    Image.new("RGB", (12, 10)).save(p2)
    slides = read_images([p1, p2], FileStore(tmp_path / "f"))
    assert [s.number for s in slides] == [1, 2]  # upload order, not filename order


@pytest.mark.parametrize(
    ("names", "code"),
    [
        (["deck.key"], "unsupported_file"),
        (["a.pdf", "b.pdf"], "validation_error"),
    ],
)
def test_validate_upload_errors(tmp_path: Path, names: list[str], code: str) -> None:
    files = []
    for n in names:
        p = tmp_path / n
        p.write_bytes(b"x")
        files.append((n, p))
    with pytest.raises(SlidexError) as exc:
        validate_upload(files)
    assert exc.value.code == code


def test_too_many_images(tmp_path: Path) -> None:
    p = tmp_path / "x.png"
    p.write_bytes(b"x")
    with pytest.raises(SlidexError) as exc:
        validate_upload([(f"{i}.png", p) for i in range(301)])
    assert exc.value.code == "too_many_slides"


def test_zip_bomb_guard(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    p = tmp_path / "big.pptx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("ppt/x.xml", "x")
    monkeypatch.setattr("slidex.intake.common.MAX_PPTX_UNCOMPRESSED", 0)
    with pytest.raises(SlidexError) as exc:
        validate_upload([("big.pptx", p)])
    assert exc.value.code == "unsupported_file"
