"""Generate the fixture decks in tests/fixtures/decks/ (T027).

    uv run python tests/fixtures/make_fixtures.py

Deterministic: re-running produces the same decks. Also copies the two demo decks to
../samples/ so judges and new users have something realistic to upload.
"""

from __future__ import annotations

import io
import math
import shutil
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
DECKS = HERE / "decks"
SAMPLES = HERE.parents[2] / "samples"

W, H = 960, 540  # 16:9 slide in points


# --------------------------------------------------------------------------- helpers


def _font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for name in ("arial.ttf", "DejaVuSans.ttf", "Helvetica.ttc"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def _png(img: Image.Image) -> bytes:
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _diagram(labels: bool, size: tuple[int, int] = (640, 360)) -> Image.Image:
    """Three replicas and a client with arrows; optionally unlabelled (a 'vague' diagram)."""
    img = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(img)
    f = _font(20)
    nodes = {
        "Client": (90, 180),
        "Replica A": (400, 70),
        "Replica B": (400, 180),
        "Replica C": (400, 290),
    }
    for name, (x, y) in nodes.items():
        d.rounded_rectangle((x - 70, y - 28, x + 70, y + 28), radius=10, outline="black", width=3)
        if labels:
            d.text((x, y), name, fill="black", font=f, anchor="mm")
    cx, cy = nodes["Client"]
    for name in ("Replica A", "Replica B", "Replica C"):
        x, y = nodes[name]
        d.line((cx + 70, cy, x - 70, y), fill="black", width=3)
        ang = math.atan2(y - cy, (x - 70) - (cx + 70))
        tip = (x - 70, y)
        for side in (-0.4, 0.4):
            d.line(
                (tip[0] - 16 * math.cos(ang + side), tip[1] - 16 * math.sin(ang + side), *tip),
                fill="black",
                width=3,
            )
    if labels:
        d.text((250, 330), "W = 2, R = 2, N = 3", fill="black", font=f, anchor="mm")
    return img


def _formula(text: str) -> Image.Image:
    img = Image.new("RGB", (900, 220), "white")
    ImageDraw.Draw(img).text((450, 110), text, fill="black", font=_font(54), anchor="mm")
    return img


# --------------------------------------------------------------------------- vague_25.pdf

VAGUE: list[str | None] = [
    "Distributed Databases\nWeek 7",
    "Why replicate?",
    "Availability\nLatency\nDurability",
    None,  # unlabelled diagram
    "R + W > N",
    "Quorums",
    "See lecture",
    "Sloppy quorum / hinted handoff",
    "Leader-based",
    "Sync vs async",
    "Replication lag!!",
    "Read-your-writes\nMonotonic reads",
    "Multi-leader → conflicts",
    "LWW",
    # deliberately wrong claim (the agent should flag it):
    "CAP: a system can be consistent, available AND partition-tolerant at the same time",
    "Vector clocks",
    "Version vectors",
    "Anti-entropy\nMerkle trees",
    "Consensus",
    "Raft: leader, term, log",
    "Paxos (skip)",
    "2PC",
    "Exam: know these",
    "Summary",
    "Questions?",
]


def make_vague_pdf(path: Path) -> None:
    doc = pymupdf.open()
    for i, text in enumerate(VAGUE, start=1):
        page = doc.new_page(width=W, height=H)
        if text is None:
            page.insert_image(pymupdf.Rect(160, 90, 800, 450), stream=_png(_diagram(labels=False)))
        else:
            size = 40 if i == 1 else 34
            page.insert_textbox(pymupdf.Rect(70, 150, 890, 460), text, fontsize=size)
        page.insert_text((W - 60, H - 24), str(i), fontsize=12)
    doc.set_metadata({"title": "Distributed Databases - Week 7"})
    doc.save(str(path), garbage=4, deflate=True)
    doc.close()


# --------------------------------------------------------------------------- clean_40.pptx

CLEAN_TOPICS = [
    ("Replication", "Copies of the same data on several machines connected by a network."),
    ("Goals", "Keep data close to users, keep working when parts fail, scale reads."),
    ("Single-leader", "All writes go to one leader; followers apply its replication log."),
    ("Synchronous followers", "The leader waits for the follower before confirming the write."),
    ("Asynchronous followers", "The leader confirms immediately; followers may lag behind."),
    ("Failover", "Detect a dead leader, choose a new one, reconfigure clients."),
    ("Replication logs", "Statement-based, write-ahead log shipping, or logical row logs."),
    ("Replication lag", "Followers can be seconds behind; reads may see stale data."),
    ("Read-your-writes", "A user always sees their own updates after they submit them."),
    ("Monotonic reads", "A user never sees time move backwards across successive reads."),
]


def make_clean_pptx(path: Path, picture: bytes) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    title_only = prs.slide_layouts[5]
    content = prs.slide_layouts[1]
    pictures = 0
    for n in range(40):
        title, body = CLEAN_TOPICS[n % len(CLEAN_TOPICS)]
        part = n // len(CLEAN_TOPICS) + 1
        if n == 12:  # a table slide
            s = prs.slides.add_slide(title_only)
            s.shapes.title.text = "Sync vs async replication"
            rows = [
                ("", "Synchronous", "Asynchronous"),
                ("Durability", "High", "Lower"),
                ("Write latency", "Higher", "Low"),
                ("Leader blocked by slow follower", "Yes", "No"),
            ]
            table = s.shapes.add_table(4, 3, Inches(1), Inches(1.8), Inches(11), Inches(3)).table
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    table.cell(r, c).text = val
        elif n in (5, 20, 33) and pictures < 3:  # picture slides
            s = prs.slides.add_slide(title_only)
            s.shapes.title.text = f"{title}: architecture"
            s.shapes.add_picture(io.BytesIO(picture), Inches(2.5), Inches(1.8), width=Inches(8))
            pictures += 1
        else:
            s = prs.slides.add_slide(content)
            s.shapes.title.text = f"{title} ({part})" if part > 1 else title
            tf = s.placeholders[1].text_frame
            tf.text = body
            tf.add_paragraph().text = "Trade-off: consistency versus latency and availability."
            for p in tf.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(24)
        s.notes_slide.notes_text_frame.text = (
            f"Speaker notes for slide {n + 1}: explain {title.lower()}."
        )
    prs.save(str(path))


# --------------------------------------------------------------------------- photos/


def make_photos(folder: Path) -> None:
    """8 'phone photos of a projector': skew, glare, blur, and EXIF rotation."""
    folder.mkdir(parents=True, exist_ok=True)
    texts = [
        "Quorum: R + W > N",
        "Leader election",
        "Replication lag",
        "Vector clocks",
        "Merkle trees",
        "Two-phase commit",
        "Raft terms",
        "Hinted handoff",
    ]
    for i, text in enumerate(texts, start=1):
        slide = Image.new("RGB", (1280, 720), (245, 245, 240))
        d = ImageDraw.Draw(slide)
        d.text((640, 300), text, fill=(30, 30, 30), font=_font(72), anchor="mm")
        d.text(
            (640, 420),
            "Lecture 7 - distributed data",
            fill=(80, 80, 80),
            font=_font(36),
            anchor="mm",
        )
        glare = Image.new("L", slide.size, 0)
        ImageDraw.Draw(glare).ellipse((700 + i * 20, 80, 1150, 400), fill=110)
        slide = Image.composite(
            Image.new("RGB", slide.size, "white"), slide, glare.filter(ImageFilter.GaussianBlur(60))
        )
        slide = slide.rotate(3 + i, expand=True, fillcolor=(60, 60, 60)).filter(
            ImageFilter.GaussianBlur(1.2)
        )
        exif = Image.Exif()
        exif[0x0112] = 6 if i % 2 else 1  # orientation: rotate 90° CW on display for odd photos
        if i % 2:
            slide = slide.transpose(Image.Transpose.ROTATE_90)  # stored sideways, EXIF fixes it
        slide.save(folder / f"photo_{i:02d}.png", exif=exif)


# --------------------------------------------------------------------------- formulas_15.pdf

FORMULAS = [
    "R + W > N",
    "W > N / 2",
    "P(fail) = 1 - (1 - p)^n",
    "lag = t_follower - t_leader",
    "quorum = floor(N/2) + 1",
    "MTTF / (MTTF + MTTR)",
    "C ≠ A under P",
    "f < N / 3",
    "T(n) = 2T(n/2) + n",
    "log2(N) hops",
    "λ = 1 / MTBF",
    "E[X] = Σ x p(x)",
    "throughput = n / t",
    "RTT = 2 × d / c",
    "Σ votes ≥ ⌈N/2⌉",
]


def make_formulas_pdf(path: Path) -> None:
    doc = pymupdf.open()
    for f in FORMULAS:
        page = doc.new_page(width=W, height=H)
        page.insert_image(pymupdf.Rect(30, 160, 930, 380), stream=_png(_formula(f)))  # image only
    doc.save(str(path), garbage=4, deflate=True)
    doc.close()


def main() -> None:
    DECKS.mkdir(parents=True, exist_ok=True)
    make_vague_pdf(DECKS / "vague_25.pdf")
    make_clean_pptx(DECKS / "clean_40.pptx", _png(_diagram(labels=True)))
    make_photos(DECKS / "photos")
    make_formulas_pdf(DECKS / "formulas_15.pdf")
    SAMPLES.mkdir(exist_ok=True)
    shutil.copy(DECKS / "vague_25.pdf", SAMPLES / "vague-lecture.pdf")
    shutil.copy(DECKS / "clean_40.pptx", SAMPLES / "clean-lecture.pptx")
    print(f"Wrote fixtures to {DECKS} and samples to {SAMPLES}")


if __name__ == "__main__":
    main()
