"""PowerPoint decks: python-pptx for text/notes/tables; LibreOffice (headless) for renders."""

from __future__ import annotations

import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.shapes.base import BaseShape
from pptx.shapes.group import GroupShape

from slidex.core.config import Settings
from slidex.core.errors import SlidexError
from slidex.core.hardware import libreoffice_path
from slidex.db.files import FileStore
from slidex.intake.common import MAX_SLIDES, SlideData
from slidex.intake.pdf import read_pdf

RENDER_TIMEOUT = 180


def _walk(shapes: list[BaseShape]) -> list[BaseShape]:
    out: list[BaseShape] = []
    for sh in shapes:
        if isinstance(sh, GroupShape):
            out.extend(_walk(list(sh.shapes)))
        else:
            out.append(sh)
    return out


def _pos(sh: BaseShape) -> tuple[int, int]:
    return (int(sh.top or 0) // 100_000, int(sh.left or 0))


def extract_pptx(path: Path) -> list[tuple[str, str, list[list[list[str]]]]]:
    """Per visible slide: (text in reading order, speaker notes, tables)."""
    try:
        pres = Presentation(str(path))
    except Exception as exc:
        raise SlidexError("corrupt_file", "The PowerPoint file could not be read.") from exc
    slides = []
    for slide in pres.slides:
        if slide._element.get("show") == "0":  # hidden slides are not exported/rendered
            continue
        texts: list[str] = []
        tables: list[list[list[str]]] = []
        for sh in sorted(_walk(list(slide.shapes)), key=_pos):
            if getattr(sh, "has_table", False):
                tables.append([[c.text.strip() for c in row.cells] for row in sh.table.rows])  # type: ignore[attr-defined]
            elif getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():  # type: ignore[attr-defined]
                texts.append(sh.text_frame.text.strip())  # type: ignore[attr-defined]
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append(("\n".join(texts), notes, tables))
    if len(slides) > MAX_SLIDES:
        raise SlidexError("too_many_slides", f"{len(slides)} slides; the limit is 300.")
    return slides


def render_pptx(path: Path, settings: Settings) -> Path:
    """Convert to PDF with LibreOffice (animations flattened to the final state)."""
    soffice = libreoffice_path(settings)
    if soffice is None:
        raise SlidexError(
            "libreoffice_missing", "Install LibreOffice or set SLIDEX_SOFFICE to soffice.exe."
        )
    outdir = Path(tempfile.mkdtemp(prefix="slidex-pptx-"))
    try:
        subprocess.run(  # noqa: S603 — fixed executable, no shell
            [
                str(soffice),
                "--headless",
                "--norestore",
                "--convert-to",
                "pdf",
                "--outdir",
                str(outdir),
                str(path),
            ],
            check=True,
            capture_output=True,
            timeout=RENDER_TIMEOUT,
        )
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
        raise SlidexError("corrupt_file", "LibreOffice could not render the presentation.") from exc
    pdf = outdir / (path.stem + ".pdf")
    if not pdf.is_file():
        raise SlidexError("corrupt_file", "LibreOffice produced no output for this presentation.")
    return pdf


def read_pptx(path: Path, files: FileStore, settings: Settings) -> list[SlideData]:
    content = extract_pptx(path)
    renders = read_pdf(render_pptx(path, settings), files, with_text=False)
    if len(renders) != len(content):
        raise SlidexError(
            "corrupt_file", f"Rendered {len(renders)} pages for {len(content)} slides."
        )
    return [
        SlideData(
            number=r.number,
            image_hash=r.image_hash,
            native_text=text,
            notes=notes,
            tables=tables,
            image_only=len(text.strip()) < 20,
        )
        for r, (text, notes, tables) in zip(renders, content, strict=True)
    ]
