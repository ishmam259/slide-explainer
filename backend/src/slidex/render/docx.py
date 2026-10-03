"""Word renderer: python-docx with headings, slide images, callouts, formula images, links."""

from __future__ import annotations

import io
import re

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from playwright.sync_api import Page

from slidex.db.files import FileStore
from slidex.llm.schemas import Block, EvidenceRef, ExplanationDocumentModel
from slidex.render import register
from slidex.render.browser import katex_css_inline, load_with_math, with_page
from slidex.render.html import LABELS

MUTE = RGBColor(0x88, 0x88, 0x88)
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def formula_images(latex: list[str]) -> list[bytes]:
    """Render each LaTeX formula to a PNG with KaTeX in one browser session."""
    if not latex:
        return []
    body = "".join(
        f'<div class="f" style="display:inline-block;padding:6px 10px;background:#fff" '
        f'data-latex="{_attr(x)}"></div><br>'
        for x in latex
    )
    head = f"<head><style>{katex_css_inline()}</style></head>"
    html = f"<!doctype html><html>{head}<body>{body}</body></html>"

    def run(page: Page) -> list[bytes]:
        page.set_viewport_size({"width": 1200, "height": 800})
        load_with_math(page, html)
        return [el.screenshot(type="png") for el in page.query_selector_all(".f")]

    return with_page(run)


def _attr(s: str) -> str:
    return s.replace("&", "&amp;").replace('"', "&quot;").replace("<", "&lt;").replace(">", "&gt;")


def _runs(p: Paragraph, text: str) -> None:
    """Minimal inline markdown: **bold**, *italic*, `code`."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*"):
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)


def _markdown(doc: DocxDocument, text: str) -> None:
    for para in re.split(r"\n\s*\n", text or ""):
        lines = [ln for ln in para.splitlines() if ln.strip()]
        if lines and all(re.match(r"^\s*([-*]|\d+\.)\s+", ln) for ln in lines):
            for ln in lines:
                style = "List Number" if re.match(r"^\s*\d+\.", ln) else "List Bullet"
                _runs(doc.add_paragraph(style=style), re.sub(r"^\s*([-*]|\d+\.)\s+", "", ln))
        elif lines:
            _runs(doc.add_paragraph(), " ".join(ln.strip() for ln in lines))


def _shade(p: Paragraph, fill: str) -> None:
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def _hyperlink(p: Paragraph, url: str, text: str) -> None:
    rid = p.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0070F3")
    rpr.append(color)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    p._p.append(link)


def _citations(doc: DocxDocument, ids: list[str], ev: dict[str, EvidenceRef]) -> None:
    refs = [ev[i] for i in ids if i in ev]
    if not refs:
        return
    p = doc.add_paragraph()
    for n, ref in enumerate(refs):
        if n:
            p.add_run(" · ").font.color.rgb = MUTE
        if ref.url:
            _hyperlink(p, ref.url, ref.label)
        else:
            r = p.add_run(ref.label)
            r.font.size = Pt(9)
            r.font.color.rgb = MUTE


@register("docx", "docx")
def render_docx(doc_model: ExplanationDocumentModel, files: FileStore) -> bytes:
    latex = [b.latex for s in doc_model.sections for b in s.blocks if b.type == "formula"]
    images = iter(formula_images(latex))
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading(doc_model.title, level=0)
    for sec in doc_model.sections:
        d.add_heading(sec.title, level=1)
        if sec.slide_image_hash and files.exists(sec.slide_image_hash):
            d.add_picture(io.BytesIO(files.read(sec.slide_image_hash)), width=Inches(6))
        for b in sec.blocks:
            _block(d, b, sec.evidence, images)
    if doc_model.sources:
        d.add_heading("Sources", level=1)
        for src in doc_model.sources:
            p = d.add_paragraph(style="List Bullet")
            prefix = (", ".join(src.authors) + ", ") if src.authors else ""
            if src.url:
                p.add_run(prefix)
                _hyperlink(p, src.url, src.title)
            else:
                p.add_run(prefix + src.title)
    if doc_model.further_reading:
        d.add_heading("Further reading", level=1)
        for src in doc_model.further_reading:
            d.add_paragraph(
                (", ".join(src.authors) + ", " if src.authors else "") + src.title,
                style="List Bullet",
            )
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _block(d: DocxDocument, b: Block, ev: dict[str, EvidenceRef], images: object) -> None:
    label = LABELS.get(b.type, "")
    if b.type == "refer_back":
        d.add_paragraph().add_run(f"{b.concept}: see Slide {b.slide}.").italic = True
        return
    if label:
        head = d.add_paragraph()
        r = head.add_run(
            label
            + (f" · confidence {round(b.confidence * 100)}%" if b.type == "reconstructed" else "")
        )
        r.bold = True
        if b.type in ("slide_says", "reconstructed", "disagreement"):
            _shade(
                head,
                {"slide_says": "F5F5F5", "reconstructed": "ECE3FA", "disagreement": "FFEFCF"}[
                    b.type
                ],
            )
    if b.type == "formula":
        png = next(images, None)  # type: ignore[call-overload]
        if png:
            d.add_picture(io.BytesIO(png))
            d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic = d.inline_shapes[-1]._inline.docPr
            pic.set("descr", b.latex)
        _markdown(d, b.steps_markdown)
    elif b.type == "slide_says":
        _markdown(d, b.text)
    else:
        _markdown(d, getattr(b, "markdown", ""))
    _citations(d, list(getattr(b, "citations", [])), ev)
